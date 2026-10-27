"""build8: after redline.py (E applied), insert purple paragraphs, accordion markers, move block, notes.
usage: python3 -I build8.py OPS.py SRC OUT"""
import sys, copy, importlib.util, docx
from docx.oxml.ns import qn
from lxml import etree
opsf, src, out = sys.argv[1:4]
s = importlib.util.spec_from_file_location('o', opsf); M = importlib.util.module_from_spec(s); s.loader.exec_module(M)
d = docx.Document(src); P = d.paragraphs
models = {'body': P[M.BODY]._p, 'list': P[M.LIST]._p}
def mk(text, kind, anchor=None):
    base = anchor if (kind == 'list' and anchor is not None) else models['body']
    q = copy.deepcopy(base)
    r0 = q.find(qn('w:r'))
    for c in list(q):
        if c.tag != qn('w:pPr'): q.remove(c)
    r = copy.deepcopy(r0)
    for c in list(r):
        if c.tag != qn('w:rPr'): r.remove(c)
    rp = r.find(qn('w:rPr'))
    if rp is None: rp = etree.SubElement(r, qn('w:rPr'))
    for tg in ('w:color', 'w:strike', 'w:b', 'w:bCs', 'w:i'):
        for e in rp.findall(qn(tg)): rp.remove(e)
    if kind == 'acc': rp.append(etree.Element(qn('w:b')))
    c = etree.SubElement(rp, qn('w:color')); c.set(qn('w:val'), '7030A0')
    t = etree.SubElement(r, qn('w:t')); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); q.append(r)
    return q
els = [p._p for p in P]
if getattr(M, 'HEAD', None):
    k, old, new = M.HEAD
    p = els[k]; r0 = p.find(qn('w:r')); t0 = r0.find(qn('w:t'))
    assert t0.text.startswith(old), t0.text
    t0.text = t0.text[len(old):]
    def run(txt, dele):
        r = copy.deepcopy(r0); [r.remove(c) for c in list(r) if c.tag != qn('w:rPr')]
        rp = r.find(qn('w:rPr'))
        if rp is None: rp = etree.SubElement(r, qn('w:rPr'))
        if dele: rp.append(etree.Element(qn('w:strike')))
        c = etree.SubElement(rp, qn('w:color')); c.set(qn('w:val'), 'C00000' if dele else '7030A0')
        t = etree.SubElement(r, qn('w:t')); t.text = txt; return r
    r0.addprevious(run(old, True)); r0.addprevious(run(new, False))
for k in sorted(M.INS, reverse=True):
    last = els[k]
    for text, kind in M.INS[k]:
        q = mk(text, kind, els[k]); last.addnext(q); last = q
if getattr(M, 'MOVE', None):
    a, b, before = M.MOVE
    blk = els[a:b + 1]; tgt = els[before]
    # move also any inserted siblings between a..b (none expected)
    for e in blk: tgt.addprevious(e)
first = [p for p in d.paragraphs if p.text.strip()][0]._p
for t in M.NOTES:
    q = mk(t, 'body'); rp = q.find(qn('w:r')).find(qn('w:rPr'))
    for e in rp.findall(qn('w:color')): rp.remove(e)
    rp.append(etree.Element(qn('w:i'))); first.addprevious(q)
d.save(out); print('ok')
# optional crop: keep notes + paragraphs CROP=(a,b) by original index, plus sectPr
if getattr(M, 'CROP', None):
    a, b = M.CROP
    d2 = docx.Document(out); body = d2.element.body
    keep = set()
    start, stop = els[a], els[b] if b < len(els) else None
    notes_n = len(M.NOTES)
    ch = list(body)
    i0 = None
    for i, e in enumerate(ch):
        if e.tag == qn('w:p') and ''.join(t.text or '' for t in e.iter(qn('w:t'))) == ''.join(t.text or '' for t in start.iter(qn('w:t'))):
            i0 = i; break
    # find stop by text too
    i1 = len(ch)
    if stop is not None:
        st = ''.join(t.text or '' for t in stop.iter(qn('w:t')))
        for i in range(i0 + 1, len(ch)):
            if ''.join(t.text or '' for t in ch[i].iter(qn('w:t'))) == st: i1 = i; break
    for i, e in enumerate(ch):
        if i < notes_n or (i0 <= i < i1) or e.tag == qn('w:sectPr'): continue
        body.remove(e)
    d2.save(out); print('cropped', i0, i1)
