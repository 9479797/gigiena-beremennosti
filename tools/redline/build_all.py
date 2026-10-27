import sys, copy, importlib.util, docx
from docx.oxml.ns import qn
from lxml import etree
opsf, src, out = sys.argv[1:4]
s = importlib.util.spec_from_file_location('o', opsf); M = importlib.util.module_from_spec(s); s.loader.exec_module(M)
d = docx.Document(src); P = d.paragraphs
body_model = P[M.BODY]._p
def mk(text, kind, anchor=None, italic=False, color='7030A0'):
    base = anchor if (kind == 'list' and anchor is not None) else body_model
    q = copy.deepcopy(base); r0 = q.find(qn('w:r'))
    for c in list(q):
        if c.tag != qn('w:pPr'): q.remove(c)
    r = copy.deepcopy(r0)
    for c in list(r):
        if c.tag != qn('w:rPr'): r.remove(c)
    rp = r.find(qn('w:rPr'))
    if rp is None: rp = etree.SubElement(r, qn('w:rPr'))
    for tg in ('w:color', 'w:strike', 'w:b', 'w:bCs', 'w:i'):
        for e in rp.findall(qn(tg)): rp.remove(e)
    if kind == 'acc': rp.append(etree.Element(qn('w:b')))
    if italic: rp.append(etree.Element(qn('w:i')))
    if color: c = etree.SubElement(rp, qn('w:color')); c.set(qn('w:val'), color)
    t = etree.SubElement(r, qn('w:t')); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); q.append(r); return q
els = [p._p for p in P]
if getattr(M, 'HEAD', None):
  k, old, new = M.HEAD
  p = els[k]; r0 = p.find(qn('w:r')); t0 = r0.find(qn('w:t'))
  assert t0.text.startswith(old), t0.text
  t0.text = t0.text[len(old):]
  def run(txt, dele):
      r = copy.deepcopy(r0); [r.remove(c) for c in list(r) if c.tag != qn('w:rPr')]
      rp = r.find(qn('w:rPr'))
      if rp is None: rp = etree.SubElement(r, qn('w:rPr'))
      if dele: rp.append(etree.Element(qn('w:strike')))
      c = etree.SubElement(rp, qn('w:color')); c.set(qn('w:val'), 'C00000' if dele else '7030A0')
      t = etree.SubElement(r, qn('w:t')); t.text = txt; return r
  r0.addprevious(run(old, True)); r0.addprevious(run(new, False))
for k in sorted(M.INS, reverse=True):
    last = els[k]
    for text, kind in M.INS[k]:
        q = mk(text, kind, els[k]); last.addnext(q); last = q
# move block: include inserted siblings up to the element before els[b+1]
if getattr(M, 'MOVE', None):
  a, b, before = M.MOVE
  blk = []; e = els[a]; endmark = els[b + 1]
  while e is not None and e is not endmark:
    blk.append(e); e = e.getnext()
  tgt = els[before]
  for e in blk: tgt.addprevious(e)
for k, notes in M.SECNOTES.items():
    for t in notes: els[k].addprevious(mk(t, 'body', italic=True, color=None))
first = [p for p in d.paragraphs if p.text.strip()][0]._p
for t in M.NOTES: first.addprevious(mk(t, 'body', italic=True, color=None))
d.save(out); print('ok')
